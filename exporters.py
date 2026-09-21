def build_print_html(sheets, sheet_w, sheet_h, kerf, colors):
    # Размер листа на печати (в мм на бумаге)
    # Ориентируемся на A4 с полями
    print_w_mm = 180
    scale = print_w_mm / sheet_w
    print_h_mm = sheet_h * scale

    parts_html = []
    parts_html.append("""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>Карта раскроя ЛДСП</title>
<style>
  @page { size: A4; margin: 10mm; }
  body { font-family: Arial, sans-serif; margin: 0; padding: 0; }
  .sheet-page {
    page-break-after: always;
    padding: 4mm;
    box-sizing: border-box;
  }
  .sheet-page:last-child { page-break-after: auto; }
  .sheet-title {
    font-size: 13pt;
    font-weight: bold;
    margin-bottom: 2mm;
  }
  .sheet-info {
    font-size: 9pt;
    color: #555;
    margin-bottom: 3mm;
  }
  .canvas {
    position: relative;
    border: 1.5px solid #000;
    background: #fff;
  }
  .part {
    position: absolute;
    border: 0.8px solid #333;
    box-sizing: border-box;
    overflow: hidden;
    font-size: 6.5pt;
    line-height: 1.15;
    padding: 1px 2px;
    text-align: center;
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: center;
    word-break: break-word;
  }
  .part .num {
    font-weight: bold;
    font-size: 7.5pt;
  }
  .part .size {
    font-size: 6pt;
    color: #222;
  }
  .part .rot {
    font-size: 5.5pt;
    color: #800;
  }
  /* Вертикальный блок текста — поворачиваем содержимое на 90° */
  .part .vertical-stack {
    position: absolute;
    left: 50%;
    top: 50%;
    transform: translate(-50%, -50%) rotate(90deg);
    transform-origin: center;
    white-space: nowrap;
    text-align: center;
    line-height: 1.15;
    font-size: 6.5pt;
  }
  .part .vertical-stack .name {
    font-weight: normal;
  }
  .part .vertical-stack .size {
    font-size: 6pt;
    color: #222;
  }
  .part .dim-top,
  .part .dim-bottom {
    position: absolute;
    left: 50%;
    transform: translateX(-50%);
    font-size: 5.5pt;
    color: #333;
    font-weight: bold;
    white-space: nowrap;
  }
  .part .dim-top { top: 1px; }
  .part .dim-bottom { bottom: 1px; }

  .part .dim-left,
  .part .dim-right {
    position: absolute;
    top: 50%;
    transform: translateY(-50%) rotate(-90deg);
    transform-origin: center;
    font-size: 5.5pt;
    color: #333;
    font-weight: bold;
    white-space: nowrap;
  }
  .part .dim-left { left: 1px; }
  .part .dim-right { right: 1px; }
  .dims-top {
    text-align: center;
    font-size: 8pt;
    font-weight: bold;
    margin-bottom: 1mm;
  }
  .dims-side {
    position: absolute;
    left: -12mm;
    top: 50%;
    transform: rotate(-90deg) translateX(50%);
    transform-origin: center;
    font-size: 8pt;
    font-weight: bold;
  }
  .legend {
    font-size: 8pt;
    margin-top: 3mm;
    border-top: 0.5px solid #999;
    padding-top: 2mm;
  }
  .legend table {
    border-collapse: collapse;
    width: 100%;
  }
  .legend td, .legend th {
    border: 0.5px solid #999;
    padding: 1px 4px;
    font-size: 7.5pt;
  }
  .legend th {
    background: #eee;
    text-align: left;
  }
  .footer {
    font-size: 7pt;
    color: #888;
    text-align: right;
    margin-top: 1mm;
  }
</style>
</head>
<body>
""")

    for i, sheet in enumerate(sheets, 1):
        parts_html.append('<div class="sheet-page">')
        parts_html.append(
            f'<div class="sheet-title">Лист {i} из {len(sheets)}</div>'
        )
        parts_html.append(
            f'<div class="sheet-info">'
            f'Размер: {sheet_w}×{sheet_h} мм · '
            f'Пил: {kerf} мм · '
            f'Заполнение: {sheet.fill_percent:.1f}% · '
            f'Деталей: {len(sheet.placed)}'
            f'</div>'
        )

        parts_html.append(
            f'<div class="dims-top">{sheet_w} мм</div>'
        )
        parts_html.append(
            f'<div class="canvas" style="width:{print_w_mm}mm; '
            f'height:{print_h_mm:.2f}mm;">'
        )

        # Детали
        for j, part in enumerate(sheet.placed):
            px = part.x * scale
            py = part.y * scale
            pw = part.w * scale
            ph = part.h * scale
            color = colors[j % len(colors)]

            # Содержимое ячейки
            inner = []

            # Определяем ориентацию: если высота больше ширины —
            # текст поворачиваем вертикально (по длинной стороне)
            vertical = part.h > part.w

            # Дублирование размеров по сторонам — только для сторон >= 500 мм
            # Длина сверху / снизу
            if part.w >= 500 and pw >= 20:
                inner.append(f'<div class="dim-top">{part.w}</div>')
                inner.append(f'<div class="dim-bottom">{part.w}</div>')

            # Высота слева / справа
            if part.h >= 500 and ph >= 20:
                inner.append(f'<div class="dim-left">{part.h}</div>')
                inner.append(f'<div class="dim-right">{part.h}</div>')

            # Название и размер — в центр, с возможным поворотом
            name = part.name.replace("<", "&lt;").replace(">", "&gt;")
            size_txt = f"{part.w}×{part.h}"
            rot_mark = ' <span class="rot">⟳</span>' if part.rotated else ""

            if vertical and pw >= 18 and ph >= 30:
                # Вертикальный блок
                inner.append(
                    '<div class="vertical-stack">'
                    f'<div class="name">{name}</div>'
                    f'<div class="size">{size_txt}{rot_mark}</div>'
                    '</div>'
                )
            else:
                # Горизонтальный (обычный)
                inner.append(f'<div class="num">{j + 1}</div>')
                if pw >= 18 and ph >= 12:
                    inner.append(f'<div>{name}</div>')
                if pw >= 15 and ph >= 9:
                    inner.append(
                        f'<div class="size">{size_txt}{rot_mark}</div>'
                    )

            parts_html.append(
                f'<div class="part" style="'
                f'left:{px:.2f}mm; top:{py:.2f}mm; '
                f'width:{pw:.2f}mm; height:{ph:.2f}mm; '
                f'background:{color};">'
                + "".join(inner) +
                '</div>'
            )

        parts_html.append('</div>')  # закрыли .canvas

        # Легенда — таблица с деталями
        parts_html.append('<div class="legend">')
        parts_html.append('<table>')
        parts_html.append(
            '<tr><th>№</th><th>Название</th><th>Размер, мм</th>'
            '<th>Позиция, мм</th><th>Поворот</th></tr>'
        )
        for j, part in enumerate(sheet.placed, 1):
            rot = "да" if part.rotated else ""
            name = part.name.replace("<", "&lt;").replace(">", "&gt;")
            parts_html.append(
                f'<tr><td>{j}</td><td>{name}</td>'
                f'<td>{part.w}×{part.h}</td>'
                f'<td>{part.x}, {part.y}</td>'
                f'<td>{rot}</td></tr>'
            )
        parts_html.append('</table>')
        parts_html.append('</div>')

        parts_html.append(
            '<div class="footer">Карта раскроя ЛДСП · '
            f'Лист {i} из {len(sheets)}</div>'
        )
        parts_html.append('</div>')  # закрыли .sheet-page

    parts_html.append("</body></html>")
    return "\n".join(parts_html)


def export_to_excel(path, sheets, sheet_w, sheet_h, kerf):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    # Удаляем дефолтный лист
    default_ws = wb.active
    wb.remove(default_ws)

    # Стили
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center",
                       wrap_text=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")

    # Сводка
    ws_sum = wb.create_sheet("Сводка")
    ws_sum.append(["Параметр", "Значение"])
    for cell in ws_sum[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    total_part_area = sum(s.used_area for s in sheets)
    total_sheet_area = sum(s.total_area for s in sheets)
    waste_pct = ((1 - total_part_area / total_sheet_area) * 100
                 if total_sheet_area else 0)

    ws_sum.append(["Размер листа",
                   f"{sheet_w}×{sheet_h} мм"])
    ws_sum.append(["Толщина пила", f"{kerf} мм"])
    ws_sum.append(["Всего листов", len(sheets)])
    ws_sum.append(["Площадь деталей", f"{total_part_area / 1e6:.3f} м²"])
    ws_sum.append(["Площадь листов", f"{total_sheet_area / 1e6:.3f} м²"])
    ws_sum.append(["Отходы", f"{waste_pct:.1f}%"])
    ws_sum.column_dimensions['A'].width = 25
    ws_sum.column_dimensions['B'].width = 25

    # Стилизуем строки сводки
    for row in ws_sum.iter_rows(min_row=2):
        for cell in row:
            cell.border = border

    # Каждый лист — отдельная вкладка со списком деталей
    for i, sheet in enumerate(sheets, 1):
        ws = wb.create_sheet(f"Лист {i}")
        ws.append(["№", "Название", "Ширина, мм", "Высота, мм",
                   "X, мм", "Y, мм", "Повёрнута"])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        for j, p in enumerate(sheet.placed, 1):
            ws.append([j, p.name, p.w, p.h, p.x, p.y,
                       "да" if p.rotated else "нет"])

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = center

        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 10
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 12

        # Строка итога
        ws.append([])
        ws.append(["", "Заполнение:", f"{sheet.fill_percent:.1f}%"])

    wb.save(path)


def export_to_txt(path, sheets, sheet_w, sheet_h, kerf):
    """Сохраняет текстовый отчёт о раскрое.

    path     — куда сохранять
    sheets   — список объектов Sheet
    sheet_w  — ширина листа (мм)
    sheet_h  — высота листа (мм)
    kerf     — толщина пила (мм)
    """
    with open(path, "w", encoding="utf-8") as f:
        f.write("Раскрой ЛДСП\n")
        f.write(f"Лист: {sheet_w}×{sheet_h} мм, пил: {kerf} мм\n")
        f.write(f"Всего листов: {len(sheets)}\n")
        total_part_area = sum(s.used_area for s in sheets)
        total_sheet_area = sum(s.total_area for s in sheets)
        f.write(f"Площадь деталей: {total_part_area / 1e6:.3f} м²\n")
        f.write(f"Площадь листов: {total_sheet_area / 1e6:.3f} м²\n")
        if total_sheet_area:
            f.write(f"Отходы: "
                    f"{(1 - total_part_area / total_sheet_area) * 100:.1f}%\n")
        f.write("=" * 50 + "\n\n")
        for i, sheet in enumerate(sheets, 1):
            f.write(f"--- ЛИСТ {i} ---\n")
            for p in sheet.placed:
                rot = " (повёрнута)" if p.rotated else ""
                f.write(f"  {p.name}: {p.w}×{p.h} мм "
                        f"@ ({p.x}, {p.y}){rot}\n")
            f.write(f"  Заполнено: {sheet.fill_percent:.1f}%\n\n")
