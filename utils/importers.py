import csv
from utils.constants import ALIASES
from utils.models import Part, Sheet, PlacedPart

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    HAS_OPENPYXL = False



# ============================================================
#  ИМПОРТ ИЗ EXCEL И CSV
# ============================================================


def detect_columns(rows):
    for i, row in enumerate(rows[:5]):
        row_lower = [str(c).strip().lower() if c is not None else "" for c in row]
        found = {}
        for key, names in ALIASES.items():
            for j, cell in enumerate(row_lower):
                if cell in names:
                    found[key] = j
                    break
        if "w" in found and "h" in found:
            return i, found
    return -1, {"name": 0, "w": 1, "h": 2, "qty": 3}


def parse_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Файл пустой")
    header_idx, col_map = detect_columns(rows)
    parts = []
    errors = []
    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue

        def get_val(key, default=None):
            j = col_map.get(key)
            if j is None or j >= len(row):
                return default
            return row[j]

        name_val = get_val("name")
        w_val = get_val("w")
        h_val = get_val("h")
        qty_val = get_val("qty", 1)
        if w_val is None or h_val is None:
            continue
        try:
            w = int(float(w_val))
            h = int(float(h_val))
            qty = int(float(qty_val)) if qty_val is not None else 1
        except (ValueError, TypeError):
            errors.append(f"Строка {i}: не удалось прочитать")
            continue
        if w <= 0 or h <= 0 or qty <= 0:
            errors.append(f"Строка {i}: некорректные значения")
            continue
        name = str(name_val).strip() if name_val is not None else f"Деталь{i}"
        parts.append(Part(name, w, h, qty))
    return parts, errors

def parse_excel_map(path):
    """Читает Excel-файл с готовой картой раскроя.
    Проходит по ВСЕМ вкладкам, ищет вкладки с колонками X и Y
    (обычно это «Лист N») и собирает из них список Sheet.

    Возвращает (list_of_sheets, errors)."""
    wb = openpyxl.load_workbook(path, data_only=True)

    # Ищем вкладки с колонками X и Y
    # Считаем, что любая вкладка, где есть заголовки X и Y — это лист раскроя.
    all_sheets = []
    errors = []

    found_any = False

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # Ищем строку с заголовками
        header_row = None
        for i, row in enumerate(rows[:10]):
            row_lower = [str(c).strip().lower() if c is not None else ""
                         for c in row]
            has_x = any("x" in cell for cell in row_lower)
            has_y = any("y" in cell for cell in row_lower)
            if has_x and has_y:
                header_row = i
                break

        if header_row is None:
            # Это не вкладка с картой — пропускаем
            continue

        # Нашли — сопоставляем колонки
        header = [str(c).strip().lower() if c is not None else ""
                  for c in rows[header_row]]

        def col_contains(name_variants):
            for j, h in enumerate(header):
                for variant in name_variants:
                    if variant in h:
                        return j
            return None

        idx_name = col_contains(["название", "имя", "деталь",
                                  "наименование", "name"])
        idx_w = col_contains(["ширина", "width"])
        idx_h = col_contains(["высота", "height"])
        idx_x = col_contains(["x"])
        idx_y = col_contains(["y"])
        idx_rot = col_contains(["повёрнута", "повернута", "rotated"])

        if None in (idx_w, idx_h, idx_x, idx_y):
            # Что-то не так с колонками — пропускаем вкладку
            errors.append(f"Вкладка «{ws.title}»: не все колонки найдены "
                          f"(нужны Ширина, Высота, X, Y)")
            continue

        sheet = Sheet(w=0, h=0)  # размеры заполним в вызывающем коде

        for i, row in enumerate(rows[header_row + 1:],
                                start=header_row + 2):
            if not row or all(c is None or str(c).strip() == ""
                              for c in row):
                continue

            def get(idx, default=None):
                if idx is None or idx >= len(row):
                    return default
                return row[idx]

            name = str(get(idx_name, f"Деталь{i}")).strip()
            try:
                w = int(float(get(idx_w)))
                h = int(float(get(idx_h)))
                x = int(float(get(idx_x)))
                y = int(float(get(idx_y)))
                rotated = False
                if idx_rot is not None:
                    rot_val = str(get(idx_rot, "")).strip().lower()
                    rotated = rot_val in ("да", "yes", "true", "1", "y")
            except (ValueError, TypeError):
            # Пропускаем итоговые строки типа "Заполнение: XX%"
            # и любые другие строки без реальных данных
            # Не пишем в errors — просто молча идём дальше
                continue

            sheet.placed.append(PlacedPart(
                name=name, x=x, y=y, w=w, h=h, rotated=rotated
            ))

        if sheet.placed:
            all_sheets.append(sheet)
            found_any = True

    if not found_any:
        raise ValueError("В файле не найдено ни одной вкладки с картой "
                         "раскроя (нет колонок X и Y)")

    return all_sheets, errors


def create_excel_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Детали"
    ws.append(["Название", "Ширина", "Высота", "Кол-во"])
    ws.append(["Боковина", 550, 720, 2])
    ws.append(["Полка", 500, 550, 4])
    ws.append(["Фасад", 596, 716, 2])
    ws.append(["Дно", 500, 550, 1])
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 10
    from openpyxl.styles import Font
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(path)


def detect_encoding(path):
    try:
        import chardet
        with open(path, "rb") as f:
            raw = f.read(100_000)
        result = chardet.detect(raw)
        enc = result.get("encoding")
        if enc:
            return enc
    except ImportError:
        chardet = None
    for enc in ("utf-8-sig", "utf-8", "cp1251", "cp866", "koi8-r"):
        try:
            with open(path, "r", encoding=enc) as f:
                f.read(1000)
            return enc
        except (UnicodeDecodeError, LookupError):
            continue
    return "utf-8"


def detect_delimiter(sample):
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        return dialect.delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in [";", ",", "\t", "|"]}
        best = max(counts, key=counts.get)
        return best if counts[best] > 0 else ";"


def parse_csv(path):
    encoding = detect_encoding(path)
    with open(path, "r", encoding=encoding, newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        delimiter = detect_delimiter(sample)
        reader = csv.reader(f, delimiter=delimiter)
        rows = [row for row in reader if any(c.strip() for c in row)]
    if not rows:
        raise ValueError("Файл пустой")
    header_idx, col_map = detect_columns(rows)
    parts = []
    errors = []
    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not row or all(c.strip() == "" for c in row):
            continue

        def get_val(key, default=None):
            j = col_map.get(key)
            if j is None or j >= len(row):
                return default
            return row[j].strip() if row[j] else default

        name_val = get_val("name")
        w_val = get_val("w")
        h_val = get_val("h")
        qty_val = get_val("qty", "1")
        if not w_val or not h_val:
            continue

        def to_num(v):
            if v is None:
                return None
            return str(v).strip().replace(",", ".").replace(" ", "")

        try:
            w = int(float(to_num(w_val)))
            h = int(float(to_num(h_val)))
            qty = int(float(to_num(qty_val))) if qty_val else 1
        except (ValueError, TypeError):
            errors.append(f"Строка {i}: не удалось прочитать")
            continue
        if w <= 0 or h <= 0 or qty <= 0:
            errors.append(f"Строка {i}: некорректные значения")
            continue
        name = str(name_val).strip() if name_val else f"Деталь{i}"
        parts.append(Part(name, w, h, qty))
    return parts, errors

